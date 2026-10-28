"""Generate the Excel batch template. Run with a Python that has openpyxl."""
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from udin_model import TEMPLATE_HEADERS  # noqa: E402

out = Path(__file__).resolve().parent / "templates" / "udin_batch_template.xlsx"
out.parent.mkdir(parents=True, exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "UDIN Batch"
ws.append(TEMPLATE_HEADERS)
for idx, header in enumerate(TEMPLATE_HEADERS, 1):
    cell = ws.cell(row=1, column=idx)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(idx)].width = max(14, min(32, len(header) + 4))

sample = {
    "category": "Certificates and Other Reports",
    "sub_category": "Net worth Certificates for Bank finances",
    "frn": "028145N - DHRUV DUA & CO",
    "signing_date": "28-10-2026",
    "client_name": "Example Private Limited",
    "pan": "AAAAA0000A",
    "description": "Net worth certificate as at 31 March 2026 for bank finance. Dry-run sample.",
    "particulars_1": "Net worth",
    "value_1": "1,25,00,000",
    "particulars_2": "Turnover",
    "value_2": "4,00,00,000",
}
ws.append([sample.get(h, "") for h in TEMPLATE_HEADERS])
ws.freeze_panes = "A2"
wb.save(out)
print("WROTE", out)
